#!/usr/bin/env python3
"""
Generate DUMMY sample data for the Prime Sales Performance Power BI model.

The September 2026 slice is calibrated so the Overview page reconciles to the
reference screenshot (director x product-group quantity/points matrix and the
15-day daily activation series). July and August 2026 are synthetic history
so month-over-month, run-rate and target analytics have something to show.

Nothing here is real e& data. Director labels are taken from the reference
screenshot; sellers (KAMs), team leaders, customers, targets and every value
are placeholders. Replace the workbook with the real extract and keep the
column names.

Outputs (relative to powerbi/):
  data/PrimeSales_SampleData.xlsx   (Excel tables: Fact_Activations, Dim_KAM,
                                     Dim_Product, Dim_Target + ReadMe sheet)
  data/csv/*.csv                    (same tables as CSV)
"""
from __future__ import annotations

import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

random.seed(20260917)

ROOT = Path(__file__).resolve().parents[1] / "data"
CSV_DIR = ROOT / "csv"

# ---------------------------------------------------------------------------
# Reference calibration (from the screenshot, September 2026, days 1-15)
# ---------------------------------------------------------------------------
DIRECTORS = ["Asim Mamoon", "Rajat Verma", "ROOPAK KUMAR"]

# (points-eligible quantity, points) per director and product group
MATRIX = {
    "Asim Mamoon":  {"Devices": (22, 94),  "Digital": (23, 133), "Fixed": (58, 614), "Mobile": (841, 1094)},
    "Rajat Verma":  {"Devices": (5, 29),   "Digital": (22, 123), "Fixed": (44, 421), "Mobile": (642, 779)},
    "ROOPAK KUMAR": {"Devices": (10, 63),  "Digital": (14, 71),  "Fixed": (53, 332), "Mobile": (613, 762)},
}
# Mobile activations with no points eligibility (VOA 2115 vs points QTY 2096)
NON_ELIGIBLE_MOBILE = {"Asim Mamoon": 8, "Rajat Verma": 6, "ROOPAK KUMAR": 5}
DAILY_SERIES = [547, 183, 218, 132, 97, 17, 91, 118, 162, 220, 137, 129, 21, 164, 130]
assert sum(DAILY_SERIES) == 2366 == sum(q for d in MATRIX.values() for q, _ in d.values()) + sum(NON_ELIGIBLE_MOBILE.values())

GSM_SHARE_OF_MOBILE = 1949 / 2115        # "Total GSM - VOA" vs Mobile VOA
MNP_SHARE_OF_GSM = 379 / 1949            # "MNP - VOA"
FNP_TOTAL_SEP = 18                        # "FNP - VOA"

# Product catalogue (placeholder names except the three fixed technologies
# visible in the reference legend). Base points are only used as weights.
PRODUCTS = [
    # Product_ID, Product Group, Product, Reporting Line, base points, weight in group
    ("P-MOB-GSM", "Mobile",  "GSM Postpaid",          "GSM",           1.0, None),
    ("P-MOB-OTH", "Mobile",  "Mobile Broadband",      "Mobile Other",  1.5, None),
    ("P-FIX-FBR", "Fixed",   "Fiber",                 "Fiber",         10.0, 0.50),
    ("P-FIX-WFA", "Fixed",   "WFA",                   "WFA",           6.0, 0.30),
    ("P-FIX-BIB", "Fixed",   "BIB_CloudPro_AI_Main",  "BIB",           15.0, 0.20),
    ("P-DIG-SVC", "Digital", "Digital Service",       "Digital",       5.5, 1.00),
    ("P-DEV-DEV", "Devices", "Device",                "Devices",       5.0, 1.00),
]
PRODUCT_BY_ID = {p[0]: p for p in PRODUCTS}

# Sellers: 6 per director, skewed productivity weights (top performer first)
SELLER_WEIGHTS = [0.30, 0.22, 0.18, 0.14, 0.10, 0.06]


def build_dim_kam():
    rows = []
    n = 0
    for d_idx, director in enumerate(DIRECTORS):
        for s in range(6):
            n += 1
            tl = f"Team Leader {d_idx * 2 + (1 if s < 3 else 2)}"
            rows.append({
                "KAM_ID": f"K{n:02d}",
                "KAM": f"KAM {n:02d}",
                "Team Leader": tl,
                "Director": director,
                "Status": "Active",
            })
    return rows


DIM_KAM = build_dim_kam()
SELLERS_BY_DIRECTOR = {d: [r for r in DIM_KAM if r["Director"] == d] for d in DIRECTORS}


def split_int(total: int, weights: list[float]) -> list[int]:
    """Split an integer across weights, preserving the exact total."""
    raw = [total * w for w in weights]
    out = [int(x) for x in raw]
    rem = total - sum(out)
    order = sorted(range(len(weights)), key=lambda i: raw[i] - out[i], reverse=True)
    for i in order[:rem]:
        out[i] += 1
    return out


def scale_points(rows: list[dict], target: float):
    """Scale row points so they sum exactly to target (2 dp)."""
    base = sum(r["Points"] for r in rows)
    if not rows or base == 0:
        return
    factor = target / base
    for r in rows:
        r["Points"] = round(r["Points"] * factor, 2)
    diff = round(target - sum(r["Points"] for r in rows), 2)
    rows[-1]["Points"] = round(rows[-1]["Points"] + diff, 2)


def make_rows_for_group(director: str, group: str, qty: int, points: float, eligible=True):
    """Create qty single-unit activation rows for one director/product group."""
    sellers = SELLERS_BY_DIRECTOR[director]
    per_seller = split_int(qty, SELLER_WEIGHTS)
    rows = []
    products = [p for p in PRODUCTS if p[1] == group]
    if group == "Mobile":
        gsm = int(round(qty * GSM_SHARE_OF_MOBILE))
        product_ids = ["P-MOB-GSM"] * gsm + ["P-MOB-OTH"] * (qty - gsm)
    else:
        counts = split_int(qty, [p[5] for p in products])
        product_ids = [pid for (pid, *_), c in zip(products, counts) for _ in range(c)]
    random.shuffle(product_ids)
    i = 0
    for seller, n in zip(sellers, per_seller):
        for _ in range(n):
            pid = product_ids[i]
            i += 1
            p = PRODUCT_BY_ID[pid]
            port_in = "None"
            base_pts = p[4]
            if pid == "P-MOB-GSM" and random.random() < MNP_SHARE_OF_GSM:
                port_in = "MNP"
                base_pts = 2.0
            rows.append({
                "KAM_ID": seller["KAM_ID"],
                "Product_ID": pid,
                "PortIn": port_in,
                "Quantity": 1,
                "Points": base_pts if eligible else 0.0,
                "PointsEligible": eligible,
                "_group": group,
            })
    if eligible:
        scale_points(rows, points)
    return rows


def assign_fnp(rows: list[dict], n: int):
    candidates = [r for r in rows if r["Product_ID"] in ("P-FIX-FBR", "P-FIX-WFA") and r["PortIn"] == "None"]
    for r in random.sample(candidates, min(n, len(candidates))):
        r["PortIn"] = "FNP"


def september_rows():
    rows = []
    for director, groups in MATRIX.items():
        for group, (qty, pts) in groups.items():
            rows.extend(make_rows_for_group(director, group, qty, pts))
        rows.extend(make_rows_for_group(director, "Mobile", NON_ELIGIBLE_MOBILE[director], 0, eligible=False))
    assign_fnp(rows, FNP_TOTAL_SEP)
    random.shuffle(rows)
    # exact daily profile from the reference
    i = 0
    for day, n in enumerate(DAILY_SERIES, start=1):
        for _ in range(n):
            rows[i]["ActivationDate"] = date(2026, 9, day)
            i += 1
    assert i == len(rows)
    return rows


def history_rows(year: int, month: int, scale: float):
    """Synthetic full month with a weekday pattern and similar mix."""
    rows = []
    for director, groups in MATRIX.items():
        for group, (qty, pts) in groups.items():
            q = int(round(qty * 2 * scale * random.uniform(0.9, 1.1)))
            p = pts / qty * q * random.uniform(0.95, 1.05)
            rows.extend(make_rows_for_group(director, group, q, round(p, 2)))
        rows.extend(make_rows_for_group(director, "Mobile", int(NON_ELIGIBLE_MOBILE[director] * 2 * scale), 0, eligible=False))
    assign_fnp(rows, int(FNP_TOTAL_SEP * 2 * scale))
    # days in month
    nxt = date(year + (month == 12), (month % 12) + 1, 1)
    days = (nxt - date(year, month, 1)).days
    weights = []
    for d in range(1, days + 1):
        wd = date(year, month, d).weekday()  # Mon=0 .. Sun=6
        w = 0.25 if wd >= 5 else 1.0          # UAE weekend Sat/Sun
        if d == 1:
            w *= 3.0                           # month-start batch, as in the reference
        weights.append(w)
    counts = split_int(len(rows), [w / sum(weights) for w in weights])
    random.shuffle(rows)
    i = 0
    for d, n in enumerate(counts, start=1):
        for _ in range(n):
            rows[i]["ActivationDate"] = date(year, month, d)
            i += 1
    return rows


def inject_exceptions(rows: list[dict]):
    """A few deliberate exceptions in August so the Data Quality page has content."""
    aug = [r for r in rows if r["ActivationDate"].month == 8]
    # 2 duplicate activation IDs (set later once IDs exist -> mark)
    for r in random.sample(aug, 2):
        r["_dup"] = True
    # 1 activation with an unmapped seller
    random.choice(aug)["KAM_ID"] = "K99"
    # 1 points-eligible row with zero points
    r = random.choice([r for r in aug if r["PointsEligible"]])
    r["Points"] = 0.0
    # 1 unmapped product
    random.choice(aug)["Product_ID"] = "P-UNKNOWN"


def finalise(rows: list[dict]):
    rows.sort(key=lambda r: (r["ActivationDate"], r["KAM_ID"]))
    out = []
    dup_pool = []
    for i, r in enumerate(rows, start=1):
        act_id = f"ACT-2026-{i:06d}"
        if r.get("_dup") and dup_pool:
            act_id = dup_pool.pop()          # reuse an earlier ID -> duplicate
        elif i % 97 == 0:
            dup_pool.append(act_id)
        out.append({
            "ActivationID": act_id,
            "ActivationDate": r["ActivationDate"].isoformat(),
            "KAM_ID": r["KAM_ID"],
            "Product_ID": r["Product_ID"],
            "Customer_PID": f"PID{random.randint(100000, 999999)}",
            "PortIn": r["PortIn"],
            "Quantity": r["Quantity"],
            "Points": r["Points"],
            "PointsEligible": r["PointsEligible"],
        })
    return out


def build_targets(fact: list[dict]):
    """Dummy monthly targets per seller: Aug actual x 1.10 for Sep; ~actual for history."""
    from collections import defaultdict
    actual = defaultdict(lambda: [0.0, 0])
    for r in fact:
        key = (r["ActivationDate"][:7], r["KAM_ID"])
        actual[key][0] += r["Points"]
        actual[key][1] += r["Quantity"]
    rows = []
    for k in DIM_KAM:
        for ym in ("2026-07", "2026-08", "2026-09"):
            pts, voa = actual.get((ym, k["KAM_ID"]), [0.0, 0])
            if ym == "2026-09":
                pts, voa = actual.get(("2026-08", k["KAM_ID"]), [0.0, 0])
                f = 1.10
            else:
                f = random.uniform(0.95, 1.20)
            rows.append({
                "MonthStart": f"{ym}-01",
                "KAM_ID": k["KAM_ID"],
                "Points Target": int(round(pts * f / 10.0)) * 10,
                "VOA Target": int(round(voa * f / 5.0)) * 5,
            })
    return rows


def write_csv(name: str, rows: list[dict]):
    CSV_DIR.mkdir(parents=True, exist_ok=True)
    with open(CSV_DIR / f"{name}.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def write_xlsx(tables: dict[str, list[dict]]):
    wb = Workbook()
    ws = wb.active
    ws.title = "ReadMe"
    lines = [
        ("PRIME SALES PERFORMANCE - DUMMY SAMPLE DATA", True),
        ("This workbook contains placeholder data only. It is NOT e& data.", True),
        ("", False),
        ("Purpose: feed the PrimeSalesPerformance.pbip Power BI project so the report opens with data.", False),
        ("As-of: activations up to 15 Sep 2026 (calibrated to the reference screenshot), plus synthetic Jul-Aug 2026 history.", False),
        ("", False),
        ("Tables (Excel Tables, same names as the Power Query queries):", True),
        ("  Fact_Activations : one row per activated line/service (grain = activation line).", False),
        ("  Dim_KAM          : seller -> Team Leader -> Director mapping.", False),
        ("  Dim_Product      : product -> product group / reporting line.", False),
        ("  Dim_Target       : monthly Points and VOA targets per seller (dummy).", False),
        ("", False),
        ("To use real data: replace the rows in each table (keep the column names and types), save, then in Power BI Desktop set the", False),
        ("DataFile parameter to this workbook's full path and Refresh.", False),
        ("", False),
        ("Definitions assumed (confirm against the official e& definitions):", True),
        ("  VOA            = volume of activations = SUM(Quantity).", False),
        ("  Points         = incentive points credited for the activation.", False),
        ("  PointsEligible = FALSE for activations that count as VOA but earn no points (explains VOA 2115 vs QTY 2096 on Mobile).", False),
        ("  PortIn         = MNP (mobile number portability), FNP (fixed number portability) or None.", False),
        ("  Deliberate August exceptions for the Data Quality page: 2 duplicate IDs, 1 unmapped KAM (K99), 1 unmapped product, 1 zero-point row.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=12 if i == 1 else 11, color="FFFFFF" if i <= 2 else "000000")
        if i <= 2:
            c.fill = PatternFill("solid", fgColor="A3122E")
    ws.column_dimensions["A"].width = 140

    for name, rows in tables.items():
        ws = wb.create_sheet(name)
        headers = list(rows[0].keys())
        ws.append(headers)
        for r in rows:
            ws.append([r[h] for h in headers])
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        ws.add_table(t)
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 4)
            ws.cell(row=1, column=i).alignment = Alignment(horizontal="center")
    wb.save(ROOT / "PrimeSales_SampleData.xlsx")


def main():
    rows = september_rows()
    rows += history_rows(2026, 8, scale=0.95)
    rows += history_rows(2026, 7, scale=0.88)
    inject_exceptions(rows)
    fact = finalise(rows)
    dim_product = [
        {"Product_ID": p[0], "Product Group": p[1], "Product": p[2], "Reporting Line": p[3], "Base Points": p[4]}
        for p in PRODUCTS
    ]
    targets = build_targets(fact)
    tables = {
        "Fact_Activations": fact,
        "Dim_KAM": DIM_KAM,
        "Dim_Product": dim_product,
        "Dim_Target": targets,
    }
    for name, t in tables.items():
        write_csv(name, t)
    write_xlsx(tables)

    # reconciliation print-out
    sep = [r for r in fact if r["ActivationDate"].startswith("2026-09")]
    kam_dir = {k["KAM_ID"]: k["Director"] for k in DIM_KAM}
    prod_grp = {p[0]: p[1] for p in PRODUCTS}
    from collections import defaultdict
    agg = defaultdict(lambda: [0, 0.0, 0])
    for r in sep:
        key = (kam_dir.get(r["KAM_ID"], "?"), prod_grp.get(r["Product_ID"], "?"))
        agg[key][0] += r["Quantity"] if r["PointsEligible"] else 0
        agg[key][1] += r["Points"]
        agg[key][2] += r["Quantity"]
    print(f"rows total={len(fact)}  Sep={len(sep)}")
    for d in DIRECTORS:
        for g in ("Devices", "Digital", "Fixed", "Mobile"):
            q, p, v = agg[(d, g)]
            tq, tp = MATRIX[d][g]
            flag = "" if (q == tq and round(p) == tp) else "  <-- MISMATCH"
            print(f"{d:14s} {g:8s} QTY={q:5d} (ref {tq:5d})  Points={p:8.1f} (ref {tp:5d})  VOA={v:5d}{flag}")
    daily = defaultdict(int)
    for r in sep:
        daily[int(r["ActivationDate"][-2:])] += r["Quantity"]
    print("daily:", [daily[d] for d in range(1, 16)])
    print("MNP Sep:", sum(r["Quantity"] for r in sep if r["PortIn"] == "MNP"),
          "FNP Sep:", sum(r["Quantity"] for r in sep if r["PortIn"] == "FNP"),
          "GSM Sep:", sum(r["Quantity"] for r in sep if r["Product_ID"] == "P-MOB-GSM"))


if __name__ == "__main__":
    main()
