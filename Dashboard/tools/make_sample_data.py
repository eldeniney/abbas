#!/usr/bin/env python3
"""Creates SYNTHETIC attendance files that exercise every edge case the Power Query logic must handle.
The files are for testing the dashboard only - they contain no real employees.

Usage: python3 make_sample_data.py [output-folder]   (default ../sample_data)
"""
import os, sys, random, datetime as dt
from openpyxl import Workbook

random.seed(42)
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "sample_data"))
os.makedirs(OUT, exist_ok=True)

HEADERS = ["Employee Code", "Agency Emp ID", "Employee Name", "Source Name", "Business Unit", "Client Name", "Cost Center", "Pay Calendar",
           "Shift Name", "Process Date", "Day", "Punch In Time", "Punch Out Time", "Total Hours", "Status", "Punch In Terminal", "Punch Out Temrinal",
           "Is Justified", "Status2"]

FIRST = ["AHMED", "MOHAMED", "SARA", "FATIMA", "OMAR", "KHALED", "NOUR", "YOUSSEF", "MARIAM", "ALI", "HUDA", "RANIA", "TAREK", "LAILA", "HASSAN", "DINA"]
LAST = ["MOHSEN", "SAYED", "SALEH", "IBRAHIM", "HASSAN", "ABDALLA", "MAHMOUD", "KAMAL", "FAWZY", "YOUSSEF", "NABIL", "SAMIR"]

CLIENTS = [("eMinds", "SMB-41984", "20085-SMB-SALES-CW41984", "SMB DIGITAL BACK OFFICE", "48 Hours Weekly"),
           ("eMinds", "SMB-41985", "20085-SMB-SALES-CW41985", "SMB FIELD SALES AM", "48 Hours Weekly"),
           ("eMinds", "CONSUMER-22011", "20090-CONS-RETAIL-CW22011", "RETAIL STORE ROTATION", "45 Hours Weekly"),
           ("Etisalat Services", "ENTERPRISE-30510", "20110-ENT-SUPPORT-CW30510", "ENTERPRISE SUPPORT DAY", "40 Hours Weekly"),
           ("Etisalat Services", "ENTERPRISE-30511", "20110-ENT-SUPPORT-CW30511", "ENTERPRISE SUPPORT NIGHT", "40 Hours Weekly")]

TERMS_IN = ["SHJ-HRB-L03-D040 LOBBY ENTRY", "DXB-HQ-GF-GATE01 ENTRY", "AUH-TWR-L01-D002 ENTRY"]
TERMS_OUT = ["SHJ-MSCP-GF-GATE04-TS EXIT EXIT", "DXB-HQ-GF-GATE01 EXIT", "AUH-TWR-L01-D002 EXIT"]

employees = []
for i in range(60):
    code = f"OB{60800 + i}"
    if i == 57:
        code = "00123"          # numeric-looking id that must stay text
    c = CLIENTS[i % len(CLIENTS)]
    employees.append({"code": code, "agency": f"PT{117600 + i}", "name": f"{random.choice(FIRST)} {random.choice(LAST)} {random.choice(LAST)}", "client": c, "night": (i % len(CLIENTS) == 4)})

def fmt_date(d):
    return d.strftime("%d-%b-%Y").upper()

def fmt_dt(d):
    return d.strftime("%d-%b-%Y %I:%M:%S %p").upper()

def hhmm(hours):
    h = int(hours); m = int(round((hours - h) * 60))
    if m == 60:
        h += 1; m = 0
    return f"{h:02d}:{m:02d}"

def month_rows(year, month, variant):
    rows = []
    days = (dt.date(year + (month == 12), (month % 12) + 1, 1) - dt.date(year, month, 1)).days
    for e in employees:
        client = e["client"]
        # employee 5 changes client / cost center / shift from the second month
        if e["code"] == "OB60805" and variant >= 1:
            client = CLIENTS[3]
        for d in range(1, days + 1):
            date = dt.date(year, month, d)
            wd = date.weekday()  # Mon=0
            status = "Present"; status2 = ""; justified = ""
            pin = pout = ""; total = ""
            if wd >= 5:
                status = random.choice(["Weekly Off", "WEEKEND"]) if variant != 2 else "Weekly Off"
            else:
                r = random.random()
                if r < 0.04:
                    status = "Absent"; justified = random.choice(["Y", "N", ""])
                elif r < 0.07:
                    status = random.choice(["Annual Leave", "Sick Leave"])
                elif r < 0.075:
                    status = "WFH"               # unrecognised -> Other
                elif r < 0.078:
                    status = ""                  # blank status
            if status == "Present":
                if e["night"]:
                    start = dt.datetime.combine(date, dt.time(22, random.randint(0, 20), random.randint(0, 59)))
                    end = start + dt.timedelta(hours=8, minutes=random.randint(0, 40))
                else:
                    start = dt.datetime.combine(date, dt.time(random.choice([7, 7, 8, 8, 9]), random.randint(0, 59), random.randint(0, 59)))
                    end = start + dt.timedelta(hours=random.choice([8, 9, 9, 9, 10]), minutes=random.randint(0, 59))
                hours = (end - start).total_seconds() / 3600
                pin, pout, total = fmt_dt(start), fmt_dt(end), hhmm(hours)
                r = random.random()
                if r < 0.03:
                    pout = ""; total = ""                     # missing punch out, blank hours
                elif r < 0.045:
                    pin = ""; total = ""                      # missing punch in
                elif r < 0.05:
                    pin = pout = ""; total = ""               # both missing
                elif r < 0.055:
                    total = "00:00"                           # zero hours recorded although punched
                elif r < 0.058:
                    total = "ab:cd"                           # invalid hours
                elif r < 0.062:
                    status2 = "Late In"
                elif r < 0.064:
                    total = hhmm(2.5)                         # short hours (review)
                elif r < 0.066:
                    total = hhmm(13.2)                        # long hours (review)
            status_out = status
            if variant == 1 and status == "Present" and random.random() < 0.3:
                status_out = random.choice(["PRESENT", "present ", " Present"])   # capitalisation / spaces
            name = e["name"] if not (variant == 2 and random.random() < 0.2) else e["name"] + "  "
            bu = client[0] if not (variant == 2 and e["code"] == "OB60850" and d < 5) else ""   # blank BU
            rows.append([e["code"], e["agency"], name, "OUTSOURCE", bu, client[1], client[2], client[4], client[3],
                         fmt_date(date), date.strftime("%A"), pin, pout, total, status_out, random.choice(TERMS_IN) if pin else "",
                         random.choice(TERMS_OUT) if pout else "", justified, status2])
    # duplicates and multi-record days
    rows.append(list(rows[10]))                                     # exact duplicate
    rows.append(list(rows[11]))                                     # exact duplicate
    second = list(rows[12]); second[11] = second[11].replace(" 0", " 1", 1) if second[11] else second[11]; second[13] = "02:00"
    rows.append(second)                                             # second (different) record same day
    # record without employee code / without date
    bad1 = list(rows[13]); bad1[0] = ""; rows.append(bad1)
    bad2 = list(rows[14]); bad2[9] = ""; rows.append(bad2)
    bad3 = list(rows[15]); bad3[9] = "31-FEB-2026"; rows.append(bad3)   # invalid date
    bad4 = list(rows[16]); bad4[11] = "7:07 AMM"; rows.append(bad4)     # invalid punch text
    return rows

def write_xlsx(path, rows, headers, title_rows=0, extra_col=False, footer=False, sheet="Attendance", extra_sheet=False):
    wb = Workbook(); ws = wb.active; ws.title = sheet
    for _ in range(title_rows):
        ws.append(["Attendance report - generated for testing", None, None])
    hdr = list(headers)
    if extra_col:
        hdr = hdr + ["Remarks"]
    ws.append(hdr)
    for r in rows:
        rr = list(r) + (["auto"] if extra_col else [])
        ws.append(rr)
    if footer:
        ws.append(["Total", None, None, None, None, None, None, None, None, None, None, None, None, None, f"{len(rows)} records"])
    if extra_sheet:
        ws2 = wb.create_sheet("Summary")
        ws2.append(["Business Unit", "Employees"]); ws2.append(["eMinds", 36]); ws2.append(["Etisalat Services", 24])
    wb.save(path)

# Month 1: standard layout, typo header, title rows above header, footer row, summary sheet
rows1 = month_rows(2026, 6, 0)
write_xlsx(os.path.join(OUT, "Attendance June 2026.xlsx"), rows1, HEADERS, title_rows=2, footer=True, extra_sheet=True)

# Month 2: different column order, corrected 'Punch Out Terminal' header, extra column, different file name pattern
order = [9, 0, 1, 2, 14, 4, 5, 6, 8, 7, 10, 11, 12, 13, 15, 16, 17, 18, 3]
hdr2 = [HEADERS[i] for i in order]
hdr2 = [h if h != "Punch Out Temrinal" else "Punch Out Terminal" for h in hdr2]
rows2 = [[r[i] for i in order] for r in month_rows(2026, 7, 1)]
write_xlsx(os.path.join(OUT, "JUL-2026_attendance_export.xlsx"), rows2, hdr2, extra_col=True, sheet="Sheet1")

# Month 3: CSV with lowercase headers and trailing spaces in names
rows3 = month_rows(2026, 8, 2)
import csv
with open(os.path.join(OUT, "attendance_2026-08.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow([h.lower() for h in HEADERS])
    for r in rows3:
        w.writerow(r)

# Noise that must be ignored
with open(os.path.join(OUT, "~$Attendance June 2026.xlsx"), "wb") as f:
    f.write(b"temp")
wb = Workbook(); ws = wb.active; ws.append(["Item", "Amount"]); ws.append(["Coffee", 12]); wb.save(os.path.join(OUT, "Unrelated expenses.xlsx"))
print("Sample data written to", OUT, "| rows:", len(rows1), len(rows2), len(rows3))
